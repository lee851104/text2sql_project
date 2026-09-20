"""把原始資料與對齊產物打包成多分頁 xlsx，供 Google 試算表一次匯入。

執行方式：

    uv run --with openpyxl python scripts/export_sheets.py

openpyxl 只有這支腳本用得到，所以沒有寫進 pyproject 的依賴，改用 --with 帶入。
資料檔在 data/ 底下、依規格 §9 不進版控，執行前需要先 `make ingest`。

產出一頁一個資料檔：原始資料 8 頁（7 個官方 CSV + units_generation.json 攤平），
對齊產物 7 頁。taipower_align 裡那 8 個「來源副本」與 data/raw 逐字相同，不另外建頁。

型別策略：能無損還原的數字才轉成數值，前導零、超過 15 位、帶 % 的一律留字串。
台電原檔同一欄會混用 3000 與 3000.5，小數位數整欄一致的欄另外套顯示格式，
讓 100.0 不會顯示成 100；位數不一致的欄沒有兩全的格式，維持通用數值顯示。
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = PROJECT_ROOT / "data" / "exports" / "台電資料_原始與對齊.xlsx"

# (分頁, 來源檔, 官方資料集)
RAW_SHEETS = (
    ("raw_units", "data/raw/units.csv", "水火力發電廠位置及機組設備（8934 / d004011）"),
    ("raw_daily", "data/raw/daily.csv", "過去電力供需資訊（19995 / d006005）"),
    ("raw_outage", "data/raw/outage.csv", "未來兩年機組大修停機排程（35393 / d006008）"),
    (
        "raw_generation_cost",
        "data/raw/generation_cost.csv",
        "各種發電方式之發電成本（10856 / d018001）",
    ),
    (
        "raw_nuclear_units",
        "data/raw/nuclear_units.csv",
        "核能發電廠位置及機組設備（10858 / d056001）",
    ),
    ("raw_re_sites", "data/raw/re_sites.csv", "再生能源各場址資料（17141 / d693002）"),
    (
        "raw_re_generation",
        "data/raw/re_generation.csv",
        "自建之各類再生能源發電量（17140 / d693001）",
    ),
)

# (分頁, 來源檔, 類型, 說明)
ALIGN_SHEETS = (
    (
        "align_crosswalk",
        "taipower_align/crosswalk.csv",
        "衍生",
        "64 個 B 欄位對機組主檔的對照與容量比（align.py）",
    ),
    (
        "align_daily_long",
        "taipower_align/daily_long.csv",
        "衍生",
        "daily 寬表轉長表，577 天 × 64 欄（final.py）",
    ),
    (
        "align_re_station_crosswalk",
        "taipower_align/re_station_crosswalk.csv",
        "衍生",
        "17141 與 17140 的發電站對齊結果（python -m align）",
    ),
    (
        "align_plants",
        "taipower_align/plants.csv",
        "人工補充",
        "34 筆電廠主檔：22 官方 + 12 補充（IPP 9、核能 3）",
    ),
    (
        "align_daily_plant_scope",
        "taipower_align/daily_plant_scope.csv",
        "人工補充",
        "64 個 B 欄位的單廠／共用授權分類",
    ),
    (
        "align_outage_plant_map",
        "taipower_align/outage_plant_map.csv",
        "人工補充",
        "138 筆歲修事件的所屬電廠對照",
    ),
    (
        "align_re_sites_supplement",
        "taipower_align/re_sites_supplement.csv",
        "人工補充",
        "場址主檔漏收的發電站，目前 1 筆：中屯風力",
    ),
)

UNITS_GENERATION_JSON = "data/raw/units_generation.json"

TOC_NOTES = (
    "原始資料 8 頁對應 data.gov.tw 的 8 個台電資料集，內容是下載後未經修改的原檔；"
    "括號內是資料集編號與資源代號。",
    "taipower_align 裡另有 8 個「來源副本」（daily、units、outage 等），"
    "與原始資料逐字相同，未重複建頁。",
    "數字欄位已存成數值可直接計算；帶 % 的字串、前導零代碼一律保留文字原樣。",
    "台電原檔的小數位本來就不一致（同一欄混用 3000 與 3000.5），"
    "整欄位數一致的欄已套格式還原，其餘顯示為通用數值。",
    "本檔由 scripts/export_sheets.py 產生，資料更新後重跑即可；它是衍生產物，不進版控。",
)

INT_PATTERN = re.compile(r"^-?(0|[1-9]\d*)$")
FLOAT_PATTERN = re.compile(r"^-?(0|[1-9]\d*)\.\d+$")


def to_cell(value: str) -> str | int | float | None:
    """只在能無損還原成原字串時才轉成數值，其餘保留文字。"""
    text = value.strip()
    if text == "":
        return None if value == "" else value
    if INT_PATTERN.match(text) and len(text.lstrip("-")) <= 15:
        return int(text)
    if FLOAT_PATTERN.match(text) and len(text.lstrip("-").replace(".", "")) <= 15:
        number = float(text)
        decimals = len(text.split(".")[1])
        if repr(number) == text or f"{number:.{decimals}f}" == text:
            return number
    return value


def column_formats(header: list[str], rows: list[list[str]]) -> dict[int, str]:
    """整欄都是數字、小數位數又一致的欄，回傳可還原字面的顯示格式。"""
    formats: dict[int, str] = {}
    for index in range(len(header)):
        decimals: Counter[int | str] = Counter()
        for row in rows:
            if index >= len(row):
                continue
            text = row[index].strip()
            if text == "":
                continue
            if INT_PATTERN.match(text):
                decimals[0] += 1
            elif FLOAT_PATTERN.match(text):
                decimals[len(text.split(".")[1])] += 1
            else:
                decimals["mixed"] += 1
                break
        if "mixed" in decimals or len(decimals) != 1:
            continue
        places = next(iter(decimals))
        if isinstance(places, int) and places > 0:
            formats[index] = "0." + "0" * places
    return formats


def add_sheet(workbook: Workbook, name: str, rows: list[list[str]]) -> tuple[int, int]:
    """建立一個分頁，回傳 (資料列數, 欄數)。"""
    sheet = workbook.create_sheet(name)
    sheet.append(rows[0])
    for row in rows[1:]:
        sheet.append([to_cell(value) for value in row])
    for index, number_format in column_formats(rows[0], rows[1:]).items():
        letter = get_column_letter(index + 1)
        for line in range(2, len(rows) + 1):
            sheet[f"{letter}{line}"].number_format = number_format
    sheet.freeze_panes = "A2"
    return len(rows) - 1, len(rows[0])


def read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def units_generation_rows(path: Path) -> list[list[str]]:
    """把即時發電量 JSON 攤成一列一機組，DateTime 併成「快照時間」欄。"""
    payload = json.loads(path.read_text(encoding="utf-8-sig"))
    columns = list(payload["aaData"][0].keys())
    snapshot = payload["DateTime"]
    rows = [["快照時間"] + columns]
    rows.extend([snapshot] + [item.get(name, "") for name in columns] for item in payload["aaData"])
    return rows


def missing_sources(root: Path) -> list[str]:
    needed = [path for _, path, _ in RAW_SHEETS]
    needed.append(UNITS_GENERATION_JSON)
    needed.extend(path for _, path, _, _ in ALIGN_SHEETS)
    return [path for path in needed if not (root / path).exists()]


def build(root: Path, out: Path) -> list[list[object]]:
    workbook = Workbook()
    workbook.remove(workbook.active)
    toc = workbook.create_sheet("_目錄")
    index: list[list[object]] = []

    for name, path, dataset in RAW_SHEETS:
        count, width = add_sheet(workbook, name, read_csv(root / path))
        index.append([name, "原始資料", path, "官方原檔逐字副本", dataset, count, width])

    count, width = add_sheet(
        workbook, "raw_units_generation", units_generation_rows(root / UNITS_GENERATION_JSON)
    )
    index.append(
        [
            "raw_units_generation",
            "原始資料",
            UNITS_GENERATION_JSON,
            "JSON 攤平：aaData 一列一機組，DateTime 併為「快照時間」欄",
            "各機組發電量即時資訊(含外購電力)（8931 / d006001）；覆寫式快照",
            count,
            width,
        ]
    )

    for name, path, kind, note in ALIGN_SHEETS:
        count, width = add_sheet(workbook, name, read_csv(root / path))
        index.append([name, "對齊產物", path, kind, note, count, width])

    toc.append(["分頁", "分類", "來源檔", "類型", "說明", "資料列數", "欄數"])
    for row in index:
        toc.append(row)
    toc.append([])
    toc.append(["註記"])
    for note in TOC_NOTES:
        toc.append(["", note])
    toc.freeze_panes = "A2"
    for letter, width in zip("ABCDEFG", (28, 10, 46, 20, 62, 10, 6), strict=True):
        toc.column_dimensions[letter].width = width

    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return index


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="輸出的 xlsx 路徑")
    args = parser.parse_args(argv)

    missing = missing_sources(PROJECT_ROOT)
    if missing:
        print("缺少來源檔，請先執行 make ingest（再視需要 make align）：", file=sys.stderr)
        for path in missing:
            print(f"  {path}", file=sys.stderr)
        return 1

    index = build(PROJECT_ROOT, args.out)
    print(f"已產生 {args.out}")
    print(f"分頁 {len(index) + 1} 個（含 _目錄）")
    for row in index:
        print(f"  {row[0]:<28} {row[1]}  {row[5]:>6} 列 × {row[6]:>2} 欄")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
